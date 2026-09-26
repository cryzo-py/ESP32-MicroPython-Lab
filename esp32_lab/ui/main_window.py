"""
Fenêtre principale (MainWindow) de l'application ESP32 MicroPython Lab
Conforme à la maquette visuelle et aux spécifications techniques
"""

from pathlib import Path
from PySide6.QtCore import QRectF, QSize, Qt, QSettings
from PySide6.QtGui import QAction, QActionGroup, QFont, QIcon, QKeySequence, QPixmap, QUndoStack
from PySide6.QtWidgets import (
    QButtonGroup,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QMainWindow,
    QMenu,
    QMessageBox,
    QPushButton,
    QSlider,
    QSplitter,
    QStackedWidget,
    QStatusBar,
    QTabWidget,
    QToolButton,
    QVBoxLayout,
    QWidget,
)

from ..app.config import AppConfig
from ..app.event_bus import get_event_bus
from ..app.i18n import tr, get_current_language, set_current_language
from ..core.models.project import ProjectModel
from ..core.project_service import ProjectService
from ..hardware.serial_manager import SerialManager
from ..hardware.uploader import ESP32Uploader
from ..simulator.engine import SimulationEngine
from .canvas.circuit_scene import CircuitScene
from .canvas.circuit_view import CircuitView
from .canvas.schematic_scene import SchematicScene
from .canvas.wire_item import WireGraphicsItem
from .dialogs.courses_dialog import CoursesDialog
from .dialogs.examples_dialog import ExamplesDialog
from .editor.code_editor import CodeEditor
from .panels.component_palette import ComponentPaletteWidget
from .panels.console_panel import ConsolePanelWidget
from .panels.hardware_status import HardwareStatusWidget
from .panels.oscilloscope_panel import OscilloscopePanel
from .panels.erc_panel import ERCPanelWidget
from ..core.erc_checker import ElectricalRulesChecker, ERCSeverity
from .panels.properties_panel import PropertiesPanelWidget
from esp32_lab.ui.panels.education.activity_panel import ActivityPanel
from .panels.repl_panel import REPLPanelWidget
from .panels.education.teacher_dashboard import TeacherDashboardWidget
from .dialogs.activity_wizard import ActivityWizardDialog
from .panels.welcome_view import WelcomeWidget


class MainWindow(QMainWindow):
    def __init__(self, app_profile: str = 'student'):
        super().__init__()
        self.setWindowTitle("ESP32 MicroPython Lab — Coder · Simuler · Expérimenter · Apprendre")
        self.resize(1400, 850)
        self.setMinimumSize(1100, 650)
        self.app_profile = app_profile

        # Core Services
        self.config = AppConfig.load()
        self.event_bus = get_event_bus()
        self.simulation_engine = SimulationEngine()
        self.serial_manager = SerialManager(self)
        self.uploader = ESP32Uploader(self.serial_manager, self)
        self.undo_stack = QUndoStack(self)
        self.menu_recent = None

        # État du projet (projet vierge par défaut)
        self.current_project: ProjectModel = ProjectService.create_empty_project()
        self.current_file_path: Path | None = None
        self._is_manually_modified: bool = False

        # Construction de l'interface
        
        self._setup_ui()
        
        try:
            from .education_services import inject_services
            inject_services(self)
        except Exception as e:
            print("Could not setup education services:", e)

        self._setup_shortcuts()
        self._connect_signals()

        # Charger le projet par défaut
        self._load_project(self.current_project)

    def _setup_ui(self):
        central_widget = QWidget(self)
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # 1. Barre de navigation supérieure (Header)
        self.header_bar = self._create_header_bar()
        main_layout.addWidget(self.header_bar)

        # 2. Pile centrale (QStackedWidget) : Page 0 = Accueil, Page 1 = Espace de travail
        self.central_stack = QStackedWidget(self)

        # Page 0 : Écran d'accueil moderne (WelcomeWidget)
        self.welcome_widget = WelcomeWidget(self)
        self.welcome_widget.new_project_requested.connect(self._on_welcome_new_project)
        self.welcome_widget.open_project_requested.connect(self._on_welcome_open_project)
        self.welcome_widget.examples_requested.connect(self._on_welcome_examples)
        self.welcome_widget.courses_requested.connect(self._on_welcome_courses)
        self.welcome_widget.recent_file_selected.connect(self._on_welcome_recent_file)
        self.welcome_widget.workspace_requested.connect(self._show_workspace_view)
        self.central_stack.addWidget(self.welcome_widget)

        # Page 1 : Séparateur principal horizontal du laboratoire
        main_splitter = QSplitter(Qt.Horizontal)
        main_splitter.setHandleWidth(6)
        main_splitter.setChildrenCollapsible(False)
        main_splitter.setOpaqueResize(True)

        # 2.1 Volet gauche : Bibliothèque de composants + Inspecteur de propriétés
        self.left_tabs = QTabWidget()
        self.left_tabs.setMinimumWidth(180)
        
        components_splitter = QSplitter(Qt.Vertical)
        components_splitter.setHandleWidth(6)
        components_splitter.setChildrenCollapsible(False)
        
        self.palette_widget = ComponentPaletteWidget(self)
        self.palette_widget.setMinimumHeight(80)
        components_splitter.addWidget(self.palette_widget)
        
        self.properties_panel = PropertiesPanelWidget(self)
        self.properties_panel.setMinimumHeight(80)
        components_splitter.addWidget(self.properties_panel)
        components_splitter.setSizes([500, 250])
        
        self.left_tabs.addTab(components_splitter, "Composants")
        
        self.activity_panel = ActivityPanel(self)
        self.left_tabs.addTab(self.activity_panel, "Activité")

        main_splitter.addWidget(self.left_tabs)

        # 2.2 Zone centrale : Éditeur + Console/REPL
        center_widget = self._create_center_widget()
        center_widget.setMinimumWidth(180)
        main_splitter.addWidget(center_widget)

        # 2.3 Volet droit : Scène de simulation 2D (Montage & Platine)
        self.simulator_widget = self._create_simulator_widget()
        self.simulator_widget.setMinimumWidth(150)
        main_splitter.addWidget(self.simulator_widget)
        self.circuit_scene.undo_stack = self.undo_stack

        # Facteurs d'étirement pour un redimensionnement souple
        main_splitter.setStretchFactor(0, 0)
        main_splitter.setStretchFactor(1, 1)
        main_splitter.setStretchFactor(2, 1)

        # Proportions équilibrées et aérées : Bibliothèque (220px), Éditeur & Outils (780px), Montage (400px)
        main_splitter.setSizes([220, 780, 400])
        self.central_stack.addWidget(main_splitter)

        main_layout.addWidget(self.central_stack, 1)

        # 3. Barre de statut
        self.status_bar = QStatusBar(self)
        self.setStatusBar(self.status_bar)

        # 4. Langue initiale et barre de menus complète (selon réglages / registre)
        from ..app.i18n import get_current_language
        saved_lang = get_current_language()
        self._set_app_language(saved_lang)

        # Thème initial (Sombre par défaut, ou selon préférences)
        settings = QSettings("ESP32Lab", "ESP32MicroPythonLab")
        saved_theme = settings.value("theme", "dark")
        self._set_app_theme(saved_theme)

        # Affichage initial selon préférences (Page d'accueil vs Espace de travail)
        show_welcome = settings.value("show_welcome_on_startup", True, type=bool)
        if show_welcome:
            self._show_welcome_view()
        else:
            self._show_workspace_view()

    def _set_app_theme(self, theme_name: str = "dark"):
        """Applique le thème (dark ou light) à l'application et synchronise tous les volets."""
        from ..app.theme import apply_theme
        from PySide6.QtWidgets import QApplication
        apply_theme(QApplication.instance(), theme_name)
        if hasattr(self, "circuit_scene") and self.circuit_scene:
            self.circuit_scene.set_theme(theme_name)
        if hasattr(self, "schematic_scene") and self.schematic_scene:
            self.schematic_scene.set_theme(theme_name)
        if hasattr(self, "code_editor") and self.code_editor:
            self.code_editor.set_theme(theme_name)
        if hasattr(self, "welcome_widget") and self.welcome_widget:
            self.welcome_widget.set_theme(theme_name)
        if hasattr(self, "console_panel") and hasattr(self.console_panel, "set_theme"):
            self.console_panel.set_theme(theme_name)
        if hasattr(self, "repl_panel") and hasattr(self.repl_panel, "set_theme"):
            self.repl_panel.set_theme(theme_name)
        if hasattr(self, "hardware_panel") and hasattr(self.hardware_panel, "set_theme"):
            self.hardware_panel.set_theme(theme_name)
        if hasattr(self, "oscilloscope") and hasattr(self.oscilloscope, "set_theme"):
            self.oscilloscope.set_theme(theme_name)
        if hasattr(self, "act_theme_dark") and hasattr(self, "act_theme_light"):
            self.act_theme_dark.setChecked(theme_name == "dark")
            self.act_theme_light.setChecked(theme_name == "light")
        settings = QSettings("ESP32Lab", "ESP32MicroPythonLab")
        settings.setValue("theme", theme_name)
        if hasattr(self, "status_bar") and self.status_bar:
            theme_label = "Clair (Light)" if theme_name == "light" else "Sombre (Dark)"
            self.status_bar.showMessage(f"Thème appliqué : {theme_label}", 3000)

    def _set_app_language(self, lang: str):
        """Définit la langue de l'interface (fr / en) et met à jour l'ensemble des composants visuels."""
        from ..app.i18n import set_current_language, tr
        set_current_language(lang)

        # 1. Mise à jour de la barre de navigation supérieure (Header)
        if hasattr(self, "app_tagline") and self.app_tagline:
            self.app_tagline.setText(tr("app_tagline"))
        if hasattr(self, "btn_nav_home") and self.btn_nav_home:
            self.btn_nav_home.setText(tr("btn_nav_home"))
            self.btn_nav_home.setToolTip(tr("btn_nav_home_tip"))
        if hasattr(self, "btn_nav_lab") and self.btn_nav_lab:
            self.btn_nav_lab.setText(tr("btn_nav_lab"))
            self.btn_nav_lab.setToolTip(tr("btn_nav_lab_tip"))

        # 2. Mise à jour des boutons d'actions d'exécution
        if hasattr(self, "btn_simulate") and self.btn_simulate:
            self.btn_simulate.setText(tr("btn_simulate"))
            self.btn_simulate.setToolTip(tr("btn_simulate_tip"))
        if hasattr(self, "btn_upload") and self.btn_upload:
            self.btn_upload.setText(tr("btn_upload"))
            self.btn_upload.setToolTip(tr("btn_upload_tip"))
        if hasattr(self, "btn_stop") and self.btn_stop:
            self.btn_stop.setText(tr("btn_stop"))
            self.btn_stop.setToolTip(tr("btn_stop_tip"))
        if hasattr(self, "btn_reset") and self.btn_reset:
            self.btn_reset.setText(tr("btn_reset"))
            self.btn_reset.setToolTip(tr("btn_reset_tip"))

        # 3. Contrôles de zoom éditeur et simulateur
        if hasattr(self, "btn_zoom_out") and self.btn_zoom_out:
            self.btn_zoom_out.setToolTip(tr("zoom_out_tip"))
        if hasattr(self, "btn_zoom_reset") and self.btn_zoom_reset:
            self.btn_zoom_reset.setToolTip(tr("zoom_reset_tip"))
        if hasattr(self, "btn_zoom_in") and self.btn_zoom_in:
            self.btn_zoom_in.setToolTip(tr("zoom_in_tip"))
        if hasattr(self, "montage_title") and self.montage_title:
            self.montage_title.setText(tr("title_montage"))
        if hasattr(self, "zoom_label") and self.zoom_label:
            self.zoom_label.setText(tr("zoom_label"))
        if hasattr(self, "btn_fit") and self.btn_fit:
            self.btn_fit.setToolTip(tr("btn_fit_tip"))

        # 4. Onglets inférieurs
        if hasattr(self, "bottom_tabs") and self.bottom_tabs:
            self.bottom_tabs.setTabText(0, tr("tab_console_repl"))
            self.bottom_tabs.setTabText(1, tr("tab_oscilloscope"))
            self.bottom_tabs.setTabText(2, tr("tab_erc"))

        # 5. Page d'accueil
        if hasattr(self, "welcome_widget") and hasattr(self.welcome_widget, "set_language"):
            self.welcome_widget.set_language(lang)

        # 6. Recréation dynamique de la barre de menus complète
        self._create_menu_bar()

        # 7. Barre d'état
        if hasattr(self, "status_bar") and self.status_bar:
            self.status_bar.showMessage(tr("lang_changed"), 3000)

    def _show_welcome_view(self):
        """Affiche la page d'accueil avec actualisation des fichiers récents."""
        if getattr(self, "app_profile", "student") == "teacher":
            if hasattr(self, "teacher_dashboard"):
                self.central_stack.setCurrentWidget(self.teacher_dashboard)
            return
        from ..app.i18n import tr
        self.welcome_widget.refresh_recent_files()
        self.central_stack.setCurrentIndex(0)
        if hasattr(self, "status_bar") and self.status_bar:
            self.status_bar.showMessage(tr("status_welcome"))

    def _show_workspace_view(self):
        """Bascule vers l'espace de travail du laboratoire."""
        from ..app.i18n import tr
        self.central_stack.setCurrentIndex(1)
        if hasattr(self, "status_bar") and self.status_bar:
            self.status_bar.showMessage(tr("status_workspace"))

    def _on_welcome_new_project(self):
        self._new_project()
        self._show_workspace_view()

    def _on_welcome_open_project(self):
        self._open_project()
        self._show_workspace_view()

    def _on_welcome_examples(self):
        self._show_workspace_view()
        self._show_examples_dialog()

    def _on_welcome_courses(self):
        self._show_workspace_view()
        self._show_courses_dialog()

    def _on_welcome_recent_file(self, file_path: str):
        self._open_recent_file(file_path)
        self._show_workspace_view()

    def _create_header_bar(self) -> QWidget:
        header_bar = QFrame()
        header_bar.setObjectName("HeaderBar")
        header_layout = QHBoxLayout(header_bar)
        header_layout.setContentsMargins(12, 4, 12, 4)
        header_layout.setSpacing(12)

        # Logo officiel et Titre
        icons_dir = Path(__file__).resolve().parent.parent / "resources" / "icons"
        icon_ico = icons_dir / "app_icon.ico"
        icon_png = icons_dir / "app_icon.png"
        icon_path = icon_ico if icon_ico.exists() else icon_png
        if icon_path.exists():
            self.setWindowIcon(QIcon(str(icon_path)))

        logo_label = QLabel()
        if icon_png.exists():
            pix = QPixmap(str(icon_png)).scaled(28, 28, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            logo_label.setPixmap(pix)
        else:
            logo_label.setText("💻")
            logo_label.setStyleSheet("font-size: 20px;")
        header_layout.addWidget(logo_label)

        title_box = QVBoxLayout()
        title_box.setSpacing(1)
        title_box.setContentsMargins(0, 0, 0, 0)
        
        app_title = QLabel("ESP32 MicroPython Lab")
        app_title.setObjectName("AppTitle")
        app_title.setMinimumWidth(210)
        title_box.addWidget(app_title)

        from ..app.i18n import tr
        self.app_tagline = QLabel(tr("app_tagline"))
        self.app_tagline.setObjectName("AppTagline")
        title_box.addWidget(self.app_tagline)

        header_layout.addLayout(title_box)
        header_layout.addSpacing(10)

        # Navigation épurée : Bouton retour Accueil et accès direct Laboratoire
        self.btn_nav_home = QPushButton(tr("btn_nav_home"))
        self.btn_nav_home.setProperty("class", "SecondaryBtn")
        self.btn_nav_home.setToolTip(tr("btn_nav_home_tip"))
        self.btn_nav_home.clicked.connect(self._show_welcome_view)
        header_layout.addWidget(self.btn_nav_home)

        self.btn_nav_lab = QPushButton(tr("btn_nav_lab"))
        self.btn_nav_lab.setProperty("class", "PrimaryBtn")
        self.btn_nav_lab.setToolTip(tr("btn_nav_lab_tip"))
        self.btn_nav_lab.clicked.connect(self._show_workspace_view)
        header_layout.addWidget(self.btn_nav_lab)
        self.btn_nav_home.setVisible(False)
        self.btn_nav_lab.setVisible(False)

        header_layout.addStretch()
        return header_bar

    def _create_center_widget(self) -> QWidget:
        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(8)

        # Splitter vertical séparant l'Éditeur en haut et Console/REPL en bas
        vert_splitter = QSplitter(Qt.Vertical)
        vert_splitter.setHandleWidth(6)
        vert_splitter.setChildrenCollapsible(False)
        vert_splitter.setOpaqueResize(True)

        # --- Partie Haute : Éditeur et barre d'actions ---
        top_box = QWidget()
        top_layout = QVBoxLayout(top_box)
        top_layout.setContentsMargins(0, 0, 0, 0)
        top_layout.setSpacing(6)

        # Onglet fichier avec contrôles de zoom typographique
        self.tabs = QTabWidget()
        self.code_editor = CodeEditor()
        self.tabs.addTab(self.code_editor, "main.py")

        # Boutons Zoom Code (- / + / 100%)
        zoom_bar = QWidget()
        zoom_bar_layout = QHBoxLayout(zoom_bar)
        zoom_bar_layout.setContentsMargins(0, 0, 6, 0)
        zoom_bar_layout.setSpacing(3)

        from ..app.i18n import tr
        self.btn_zoom_out = QPushButton("A-")
        self.btn_zoom_out.setToolTip(tr("zoom_out_tip"))
        self.btn_zoom_out.setFixedSize(26, 22)
        self.btn_zoom_out.setStyleSheet("QPushButton { font-weight: bold; font-size: 10px; padding: 0; }")
        self.btn_zoom_out.clicked.connect(lambda: self.code_editor.zoom_out())
        zoom_bar_layout.addWidget(self.btn_zoom_out)

        self.btn_zoom_reset = QPushButton("100%")
        self.btn_zoom_reset.setToolTip(tr("zoom_reset_tip"))
        self.btn_zoom_reset.setFixedHeight(22)
        self.btn_zoom_reset.setStyleSheet("QPushButton { font-size: 10px; padding: 0 4px; }")
        self.btn_zoom_reset.clicked.connect(lambda: self.code_editor.zoom_reset())
        zoom_bar_layout.addWidget(self.btn_zoom_reset)

        self.btn_zoom_in = QPushButton("A+")
        self.btn_zoom_in.setToolTip(tr("zoom_in_tip"))
        self.btn_zoom_in.setFixedSize(26, 22)
        self.btn_zoom_in.setStyleSheet("QPushButton { font-weight: bold; font-size: 10px; padding: 0; }")
        self.btn_zoom_in.clicked.connect(lambda: self.code_editor.zoom_in())
        zoom_bar_layout.addWidget(self.btn_zoom_in)

        self.tabs.setCornerWidget(zoom_bar, Qt.TopRightCorner)
        top_layout.addWidget(self.tabs, 1)

        # Barre d'actions d'exécution (Simuler, Téléverser, Stop, Reset)
        action_bar = QHBoxLayout()
        action_bar.setContentsMargins(0, 0, 0, 0)
        action_bar.setSpacing(8)

        self.btn_simulate = QPushButton(tr("btn_simulate"))
        self.btn_simulate.setObjectName("BtnSimulate")
        self.btn_simulate.setProperty("class", "PrimaryBtn")
        self.btn_simulate.setToolTip(tr("btn_simulate_tip"))
        self.btn_simulate.clicked.connect(self._start_simulation)
        action_bar.addWidget(self.btn_simulate)

        self.btn_upload = QPushButton(tr("btn_upload"))
        self.btn_upload.setObjectName("BtnUpload")
        self.btn_upload.setProperty("class", "UploadBtn")
        self.btn_upload.setToolTip(tr("btn_upload_tip"))
        self.btn_upload.clicked.connect(self._upload_to_esp32)
        action_bar.addWidget(self.btn_upload)

        self.btn_stop = QPushButton(tr("btn_stop"))
        self.btn_stop.setObjectName("BtnStop")
        self.btn_stop.setProperty("class", "StopBtn")
        self.btn_stop.setEnabled(False)
        self.btn_stop.setToolTip(tr("btn_stop_tip"))
        self.btn_stop.clicked.connect(self._stop_simulation)
        action_bar.addWidget(self.btn_stop)

        self.btn_reset = QPushButton(tr("btn_reset"))
        self.btn_reset.setObjectName("BtnReset")
        self.btn_reset.setProperty("class", "SecondaryBtn")
        self.btn_reset.setToolTip(tr("btn_reset_tip"))
        self.btn_reset.clicked.connect(self._reset_simulation)
        action_bar.addWidget(self.btn_reset)

        action_bar.addStretch()
        top_box.setMinimumHeight(100)
        top_layout.addLayout(action_bar)
        vert_splitter.addWidget(top_box)

        # --- Partie Basse : Onglets Outils (Console & REPL / Oscilloscope) ---
        self.bottom_tabs = QTabWidget()
        self.bottom_tabs.setObjectName("BottomToolsTabs")
        self.bottom_tabs.setMinimumHeight(60)
        self.bottom_tabs.setMinimumWidth(150)

        # Onglet 1 : Console & REPL & Carte physique
        tools_container = QWidget()
        tools_container.setMinimumHeight(50)
        tools_layout = QHBoxLayout(tools_container)
        tools_layout.setContentsMargins(0, 0, 0, 0)
        tools_layout.setSpacing(8)

        bottom_splitter = QSplitter(Qt.Horizontal)
        bottom_splitter.setHandleWidth(6)
        bottom_splitter.setChildrenCollapsible(False)
        bottom_splitter.setOpaqueResize(True)
        self.console_panel = ConsolePanelWidget()
        self.console_panel.setMinimumWidth(60)
        bottom_splitter.addWidget(self.console_panel)

        self.repl_panel = REPLPanelWidget()
        self.repl_panel.setMinimumWidth(60)
        bottom_splitter.addWidget(self.repl_panel)

        self.hardware_panel = HardwareStatusWidget()
        self.hardware_panel.setMinimumWidth(60)
        bottom_splitter.addWidget(self.hardware_panel)

        bottom_splitter.setSizes([260, 260, 180])
        tools_layout.addWidget(bottom_splitter)
        self.bottom_tabs.addTab(tools_container, "📟 Console Série & REPL")

        # Onglet 2 : Oscilloscope & Analyseur Logique
        self.oscilloscope = OscilloscopePanel(self)
        self.oscilloscope.setMinimumHeight(50)
        self.oscilloscope.setMinimumWidth(120)
        self.bottom_tabs.addTab(self.oscilloscope, "📊 Oscilloscope & Analyseur Logique")

        # Onglet 3 : Diagnostic des Règles Électriques (ERC)
        self.erc_panel = ERCPanelWidget(self)
        self.erc_panel.setMinimumHeight(50)
        self.erc_panel.setMinimumWidth(120)
        self.bottom_tabs.addTab(self.erc_panel, "🛡️ Diagnostic Électrique (ERC)")
        self.erc_panel.btn_refresh.clicked.connect(self._run_erc_check)

        vert_splitter.addWidget(self.bottom_tabs)

        # Facteurs d'étirement pour l'éditeur (1) et la console (0)
        vert_splitter.setStretchFactor(0, 1)
        vert_splitter.setStretchFactor(1, 0)

        # Proportions équilibrées : Éditeur (540px), Outils bas (220px)
        vert_splitter.setSizes([540, 220])
        layout.addWidget(vert_splitter)
        return container

    def _create_simulator_widget(self) -> QWidget:
        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(6)

        # Barre d'outils du simulateur
        sim_header = QHBoxLayout()
        sim_header.setContentsMargins(0, 0, 0, 0)

        # Titre épuré du volet Montage & Platine physique
        from ..app.i18n import tr
        self.montage_title = QLabel(tr("title_montage"))
        self.montage_title.setStyleSheet("color: #38bdf8; font-weight: bold; font-size: 12px; padding-left: 2px;")
        sim_header.addWidget(self.montage_title)

        sim_header.addStretch()

        # Contrôle du Zoom
        self.zoom_label = QLabel(tr("zoom_label"))
        self.zoom_label.setObjectName("ZoomLabel")
        sim_header.addWidget(self.zoom_label)

        self.zoom_slider = QSlider(Qt.Horizontal)
        self.zoom_slider.setRange(50, 200)
        self.zoom_slider.setValue(100)
        self.zoom_slider.setFixedWidth(90)
        self.zoom_slider.valueChanged.connect(self._on_zoom_slider_changed)
        sim_header.addWidget(self.zoom_slider)

        self.zoom_val_label = QLabel("100%")
        self.zoom_val_label.setObjectName("ZoomValLabel")
        sim_header.addWidget(self.zoom_val_label)

        self.btn_fit = QPushButton("⛶")
        self.btn_fit.setToolTip(tr("btn_fit_tip"))
        self.btn_fit.setFixedWidth(28)
        self.btn_fit.clicked.connect(self._fit_view)
        sim_header.addWidget(self.btn_fit)

        layout.addLayout(sim_header)

        # Vue graphique 2D du circuit (Maquette Réaliste et Schéma Électrique CEI)
        self.circuit_scene = CircuitScene(self)
        self.schematic_scene = SchematicScene(self)
        self.circuit_view = CircuitView(self.circuit_scene, self)
        self.circuit_view.setMinimumWidth(100)
        layout.addWidget(self.circuit_view, 1)


        return container

    def _setup_shortcuts(self):
        # Raccourcis clavier
        self.act_save_shortcut = QAction(self)
        self.act_save_shortcut.setShortcut(QKeySequence("Ctrl+S"))
        self.act_save_shortcut.triggered.connect(self._save_project)
        self.addAction(self.act_save_shortcut)

        self.act_open_shortcut = QAction(self)
        self.act_open_shortcut.setShortcut(QKeySequence("Ctrl+O"))
        self.act_open_shortcut.triggered.connect(self._open_project)
        self.addAction(self.act_open_shortcut)

        self.act_new_shortcut = QAction(self)
        self.act_new_shortcut.setShortcut(QKeySequence("Ctrl+N"))
        self.act_new_shortcut.triggered.connect(self._new_project)
        self.addAction(self.act_new_shortcut)

        self.act_save_as_shortcut = QAction(self)
        self.act_save_as_shortcut.setShortcut(QKeySequence("Ctrl+Shift+S"))
        self.act_save_as_shortcut.triggered.connect(self._save_project_as)
        self.addAction(self.act_save_as_shortcut)

        self.act_sim = QAction(self)
        self.act_sim.setShortcut(Qt.Key_F5)
        self.act_sim.triggered.connect(self._start_simulation)
        self.addAction(self.act_sim)

        self.act_stop = QAction(self)
        self.act_stop.setShortcut(Qt.Key_F6)
        self.act_stop.triggered.connect(self._stop_simulation)
        self.addAction(self.act_stop)

    def _connect_signals(self):
        # Événements Palette vers Scène et Propriétés
        self.palette_widget.component_selected.connect(self._on_palette_component_selected)

        # Événements Matériel
        self.hardware_panel.port_selected.connect(self._on_port_selected)
        if hasattr(self, "console_panel") and hasattr(self.console_panel, "baudrate_combo"):
            self.console_panel.baudrate_combo.currentTextChanged.connect(self._on_baudrate_changed)
        self.hardware_panel.rescan_requested.connect(self.serial_manager.scan_ports)

        # Événements GPIO pour le badge d'état
        self.event_bus.gpio_changed.connect(self._on_gpio_status_update)
        self.event_bus.project_modified.connect(lambda: setattr(self, '_is_manually_modified', True))
        self.event_bus.simulation_stopped.connect(self._on_simulation_stopped_ui)
        self.event_bus.session_locked.connect(self._on_session_locked)

        # Connexions Inspecteur de propriétés
        self.circuit_scene.selectionChanged.connect(self._on_scene_selection_changed)
        self.properties_panel.property_changed.connect(self._on_property_changed)

    def _on_palette_component_selected(self, comp_type: str):
        """Affiche les propriétés du composant sélectionné dans la bibliothèque."""
        ctype = comp_type.lower()
        defaults = {
            "led": {"color": "red"},
            "resistor": {"value": 220},
            "potentiometer": {"raw_value": 2048},
            "dht22": {"temperature": 24.5, "humidity": 55.0},
            "hcsr04": {"distance_cm": 25.0},
            "servo": {"angle": 90.0},
            "oled": {"width": 128, "height": 64},
            "lcd": {},
            "button": {},
            "buzzer": {},
            "relay": {},
            "esp32": {},
        }
        props = dict(defaults.get(ctype, {}))
        # Si un composant est déjà sélectionné sur la platine, le garder
        selected = [it for it in self.circuit_scene.selectedItems() if getattr(it, "component_id", "").startswith(ctype)]
        if selected:
            self._on_scene_selection_changed()
            return

        self.properties_panel.inspect_component(f"Nouveau {ctype}", ctype, props)
        self.status_bar.showMessage(f"Composant '{comp_type}' sélectionné. Glissez-le sur la platine pour le placer.", 4000)

    def _show_courses_dialog(self):
        self.current_project.set_main_code(self.code_editor.toPlainText())
        dlg = CoursesDialog(self.current_project, self)
        dlg.lesson_loaded.connect(self._on_course_lesson_loaded)
        dlg.exec()

    def _on_course_lesson_loaded(self, p):
        if not self._maybe_save_changes("charger cet exercice"):
            return
        self.current_file_path = None
        self._load_project(p)
        self._show_workspace_view()

    def _show_examples_dialog(self):
        dlg = ExamplesDialog(self)
        dlg.example_loaded.connect(self._on_example_project_loaded)
        dlg.exec()

    def _on_example_project_loaded(self, p):
        if not self._maybe_save_changes("charger cet exemple"):
            return
        self.current_file_path = None
        self._load_project(p)
        self._show_workspace_view()

    def _on_scene_selection_changed(self):
        items = self.circuit_scene.selectedItems()
        if not items:
            self.properties_panel.clear_selection()
            return

        # Trouver le composant ciblé (en remontant les ancres de broches si nécessaire)
        target_item = None
        for it in items:
            curr = it
            while curr:
                if hasattr(curr, "component_id") and getattr(curr, "component_id") not in ("breadboard_1", "esp32"):
                    target_item = curr
                    break
                curr = curr.parentItem()
            if target_item:
                break

        # Si aucun composant secondaire, vérifier si l'ESP32 est sélectionné
        if not target_item:
            for it in items:
                curr = it
                while curr:
                    if hasattr(curr, "component_id") and getattr(curr, "component_id") == "esp32":
                        target_item = curr
                        break
                    curr = curr.parentItem()
                if target_item:
                    break

        if not target_item:
            self.properties_panel.clear_selection()
            return

        comp_id = getattr(target_item, "component_id", None)
        if not comp_id or comp_id == "breadboard_1":
            self.properties_panel.clear_selection()
            return

        # Déterminer le type du composant
        comp_type = getattr(target_item, "_comp_type", "")
        if not comp_type:
            cls_name = target_item.__class__.__name__.lower()
            for t in ("resistor", "led", "rgb_led", "ldr", "pir", "joystick", "switch", "button", "potentiometer", "dht22", "oled", "hcsr04", "lcd", "relay", "servo", "buzzer", "neopixel", "board", "esp32"):
                if t in comp_id.lower() or t in cls_name:
                    comp_type = "esp32" if t == "board" else t
                    break
        if not comp_type:
            comp_type = "component"

        # Extraire les propriétés réelles directement de l'item graphique
        props = {}
        if hasattr(target_item, "resistance_value"):
            props["value"] = target_item.resistance_value
        if hasattr(target_item, "color_name"):
            props["color"] = target_item.color_name
        if hasattr(target_item, "raw_value"):
            props["raw_value"] = target_item.raw_value
        if hasattr(target_item, "temperature"):
            props["temperature"] = target_item.temperature
        if hasattr(target_item, "humidity"):
            props["humidity"] = target_item.humidity
        if hasattr(target_item, "distance_cm"):
            props["distance_cm"] = target_item.distance_cm
        if hasattr(target_item, "angle"):
            props["angle"] = target_item.angle
        if hasattr(target_item, "lux"):
            props["lux"] = target_item.lux
        if hasattr(target_item, "val_r"):
            props["r"] = target_item.val_r
            props["g"] = target_item.val_g
            props["b"] = target_item.val_b
        if hasattr(target_item, "motion_detected"):
            props["motion_detected"] = target_item.motion_detected
        if hasattr(target_item, "x_val"):
            props["x"] = target_item.x_val
            props["y"] = target_item.y_val
        if hasattr(target_item, "is_on") and comp_type == "switch":
            props["is_on"] = target_item.is_on

        # Synchroniser ou créer dans current_project.components
        comp = next((c for c in self.current_project.components if c.id == comp_id), None)
        if comp:
            for k, v in comp.properties.items():
                if k not in props and v is not None:
                    props[k] = v
            comp.properties.update(props)
        else:
            from ..core.models.component import ComponentModel
            new_comp = ComponentModel(
                id=comp_id,
                type=comp_type,
                name=comp_type.capitalize(),
                x=target_item.scenePos().x(),
                y=target_item.scenePos().y(),
                properties=dict(props)
            )
            self.current_project.components.append(new_comp)

        self.properties_panel.inspect_component(comp_id, comp_type, props)

    def _on_property_changed(self, comp_id: str, updated_props: dict):
        self._is_manually_modified = True
        comp = next((c for c in self.current_project.components if c.id == comp_id), None)
        if comp:
            comp.properties.update(updated_props)
        else:
            # Si le composant n'est pas encore dans current_project.components (ex: ajouté dynamiquement)
            from ..core.models.component import ComponentModel
            item = self.circuit_scene.component_items.get(comp_id)
            if item:
                comp_type = getattr(self.properties_panel, "current_component_type", "component")
                new_comp = ComponentModel(
                    id=comp_id,
                    type=comp_type,
                    name=comp_type.capitalize(),
                    x=item.scenePos().x(),
                    y=item.scenePos().y(),
                    properties=dict(updated_props)
                )
                self.current_project.components.append(new_comp)

        item = self.circuit_scene.component_items.get(comp_id)
        if item:
            if hasattr(item, "set_values") and "temperature" in updated_props:
                item.set_values(updated_props["temperature"], updated_props["humidity"])
            elif hasattr(item, "color_name") and "color" in updated_props:
                item.color_name = updated_props["color"]
                item.update()
            elif hasattr(item, "resistance_value") and "value" in updated_props:
                item.set_value(int(updated_props["value"]))
            elif hasattr(item, "set_value") and "value" in updated_props:
                item.set_value(int(updated_props["value"]))
            elif hasattr(item, "raw_value") and "raw_value" in updated_props:
                item.raw_value = int(updated_props["raw_value"])
                item.update()
            elif hasattr(item, "distance_cm") and "distance_cm" in updated_props:
                item.distance_cm = float(updated_props["distance_cm"])
                item.update()
            elif hasattr(item, "set_angle") and "angle" in updated_props:
                item.set_angle(float(updated_props["angle"]))
            elif hasattr(item, "lux") and "lux" in updated_props:
                item.lux = float(updated_props["lux"])
                item.update()
            elif hasattr(item, "set_color_rgb") and ("r" in updated_props or "g" in updated_props or "b" in updated_props):
                item.set_color_rgb(
                    updated_props.get("r", getattr(item, "val_r", 255)),
                    updated_props.get("g", getattr(item, "val_g", 0)),
                    updated_props.get("b", getattr(item, "val_b", 0))
                )
            elif hasattr(item, "trigger_motion") and updated_props.get("motion_detected"):
                item.trigger_motion()
            elif hasattr(item, "set_position") and ("x" in updated_props or "y" in updated_props):
                item.set_position(
                    updated_props.get("x", getattr(item, "x_val", 2048)),
                    updated_props.get("y", getattr(item, "y_val", 2048)),
                    updated_props.get("sw", False)
                )
            elif hasattr(item, "set_state") and "is_on" in updated_props:
                item.set_state(bool(updated_props["is_on"]))
            else:
                item.update()
        self.schematic_scene.load_project_schematic(self.current_project)

    def _load_project(self, project: ProjectModel):
        self.current_project = project
        if hasattr(self, 'activity_panel'):
            self.activity_panel.load_project(project)
            is_act = bool(project.pedagogy_profile)
            is_teacher = getattr(self, "app_profile", "student") == "teacher"
            disable_menus = is_act and not is_teacher
            
            # Disable menus to prevent cheating during exams (unless teacher)
            if hasattr(self, "menu_examples"):
                self.menu_examples.setEnabled(not disable_menus)
            if hasattr(self, "menu_courses"):
                self.menu_courses.setEnabled(not disable_menus)
            if hasattr(self, "menu_course"):
                self.menu_course.setEnabled(not disable_menus)
            
            if hasattr(self, "left_tabs"):
                self.left_tabs.setTabVisible(1, is_act)
                # Ensure the Composants tab is always visible so students can build during exams
                self.left_tabs.setTabVisible(0, True)
                if is_act:
                    self.left_tabs.setCurrentIndex(1)
                else:
                    self.left_tabs.setCurrentIndex(0)

        if hasattr(self, "undo_stack") and self.undo_stack:
            self.undo_stack.clear()
        self.code_editor.setPlainText(project.get_main_code())
        self.circuit_scene.load_project_circuit(project)
        self.schematic_scene.load_project_schematic(project)
        self._mark_as_clean()
        if self.current_file_path:
            self.setWindowTitle(f"ESP32 MicroPython Lab — {self.current_file_path.name}")
        elif project.name and project.name != "Nouveau Montage":
            self.setWindowTitle(f"ESP32 MicroPython Lab — {project.name}")
        else:
            self.setWindowTitle("ESP32 MicroPython Lab")
        self.circuit_view.centerOn(220, 160)

    def _show_simulation_view(self):
        """Affiche la maquette physique réaliste du laboratoire."""
        self.circuit_view.setScene(self.circuit_scene)

    def _show_schematic_view(self):
        """Affiche le schéma électronique normalisé (CEI / IEEE)."""
        self.schematic_scene.load_project_schematic(self.current_project)
        self.circuit_view.setScene(self.schematic_scene)

    def _toggle_wire_routing_mode(self, checked: bool):
        """Bascule entre câblage souple réaliste et câblage orthogonal 90°."""
        self.circuit_scene.set_orthogonal_routing(checked)
        if checked:
            if hasattr(self, "btn_route_mode"):
                self.btn_route_mode.setText("📐 Câblage Orthogonal 90°")
            self.status_bar.showMessage("Mode de câblage : Orthogonal 90° (Angles droits industriels)", 3000)
        else:
            if hasattr(self, "btn_route_mode"):
                self.btn_route_mode.setText("〰️ Câblage Souple")
            self.status_bar.showMessage("Mode de câblage : Câblage souple réaliste (Courbes de Bézier)", 3000)

    def _on_add_component(self, ctype: str):
        if hasattr(self, "undo_stack") and self.undo_stack:
            from .canvas.undo_commands import AddComponentCommand
            self.undo_stack.push(AddComponentCommand(self, ctype))
        else:
            from ..core.models.component import ComponentModel, ComponentPin
            new_comp = ComponentModel(
                type=ctype,
                name=ctype.capitalize(),
                x=350,
                y=200,
                properties={"color": "red" if ctype == "led" else None}
            )
            self.current_project.components.append(new_comp)
            self.circuit_scene.load_project_circuit(self.current_project)
            self.schematic_scene.load_project_schematic(self.current_project)
            self.status_bar.showMessage(f"Composant '{ctype}' ajouté à la platine.", 3000)

    def _run_erc_check(self) -> list:
        """Exécute l'analyse des règles électriques sur le circuit actuel."""
        self._sync_scene_to_project()
        checker = ElectricalRulesChecker()

        comps_data = [
            {"id": c.id, "type": c.type, "properties": dict(c.properties)}
            for c in self.current_project.components
        ]
        conns_data = [
            {
                "from_component": c.from_component,
                "from_pin": c.from_pin,
                "to_component": c.to_component,
                "to_pin": c.to_pin,
            }
            for c in self.current_project.connections
        ]

        # Récupérer l'état actuel des GPIO simulés si disponible
        gpio_states = {}
        if hasattr(self, "simulation_engine") and hasattr(self.simulation_engine, "gpio_manager"):
            mgr = self.simulation_engine.gpio_manager
            for p, mode in getattr(mgr, "_modes", {}).items():
                val = mgr.read(p)
                m_str = "OUT" if int(mode) == 3 else "IN"
                gpio_states[p] = {"mode": m_str, "value": val}

        violations = checker.analyze(
            comps_data,
            conns_data,
            topology=self.circuit_scene.topology,
            gpio_states=gpio_states,
        )
        self.event_bus.erc_violations_updated.emit(violations)

        errors = [v for v in violations if v.severity == ERCSeverity.ERROR]
        if errors:
            self.bottom_tabs.setCurrentWidget(self.erc_panel)
            self.status_bar.showMessage(f"⚠️ Alerte ERC : {errors[0].title}", 5000)
        return violations

    def _start_simulation(self):
        # Vérification des règles électriques avant exécution
        violations = self._run_erc_check()
        severe_errors = [v for v in violations if v.severity == ERCSeverity.ERROR]
        if severe_errors:
            reply = QMessageBox.warning(
                self,
                "Erreur Électrique Détectée (ERC)",
                f"<b>{severe_errors[0].title}</b><br><br>"
                f"{severe_errors[0].message}<br><br>"
                f"💡 <i>{severe_errors[0].recommendation}</i><br><br>"
                "Voulez-vous quand même forcer le démarrage de la simulation ?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if reply != QMessageBox.Yes:
                return

        code = self.code_editor.toPlainText()
        self.current_project.set_main_code(code)
        self.code_editor.clear_error_highlight()
        self.btn_simulate.setEnabled(False)
        self.btn_stop.setEnabled(True)
        self.simulation_engine.start(code)
        self.status_bar.showMessage("Simulation en cours d'exécution...")

    def _stop_simulation(self):
        self.simulation_engine.stop()
        self.btn_simulate.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self.status_bar.showMessage("Simulation arrêtée.")

    def _reset_simulation(self):
        self.simulation_engine.reset()
        self.btn_simulate.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self.status_bar.showMessage("Simulation réinitialisée.")

    def _on_toggle_network_bridge(self, checked: bool):
        from ..simulator.modules.network import set_host_bridge_enabled
        set_host_bridge_enabled(checked)
        settings = QSettings("ESP32Lab", "ESP32MicroPythonLab")
        settings.setValue("use_host_network_bridge", checked)
        mode_str = "🌐 Passerelle Réseau Réelle (Hôte PC) activée." if checked else "🔒 Mode Réseau Sandbox (Simulé) activé."
        self.status_bar.showMessage(mode_str, 4000)

    def _on_simulation_stopped_ui(self):
        self.btn_simulate.setEnabled(True)
        self.btn_stop.setEnabled(False)

    def _on_gpio_status_update(self, pin: int, value: int):
        if pin == 2:
            state_text = "🟢 Allumée" if value else "⚪ Éteinte"
    
    def _upload_to_esp32(self):
        code = self.code_editor.toPlainText()
        if not self.serial_manager.current_port_name:
            # Demander de scanner ou sélectionner un port
            ports = self.serial_manager.scan_ports()
            if not ports:
                QMessageBox.warning(
                    self,
                    "Aucune carte détectée",
                    "Veuillez brancher une carte ESP32 via le câble USB.\n"
                    "Si elle est branchée, vérifiez vos pilotes USB-Série (CH340 / CP210x)."
                )
                return
        self.uploader.upload_code(code, "main.py")

    def _on_port_selected(self, port_name: str):
        if port_name:
            baud_str = self.console_panel.baudrate_combo.currentText()
            baudrate = int(baud_str) if baud_str.isdigit() else 115200
            self.serial_manager.connect_port(port_name, baudrate)

    def _on_baudrate_changed(self, baud_str: str):
        # Reconnect if already connected to a port
        if self.serial_manager.is_connected() and self.serial_manager.serial_port:
            port_name = self.serial_manager.serial_port.port
            self.serial_manager.disconnect_port()
            self._on_port_selected(port_name)

    def _on_zoom_slider_changed(self, value: int):
        self.zoom_val_label.setText(f"{value}%")
        scale = value / 100.0
        self.circuit_view.resetTransform()
        self.circuit_view.scale(scale, scale)

    def _fit_view(self):
        # Calculer le rectangle englobant des composants physiques réels
        visible_rect = QRectF()
        for item in self.circuit_scene.items():
            if item != self.circuit_scene.usb_cable_item and item.isVisible():
                if hasattr(item, "component_id") or isinstance(item, WireGraphicsItem):
                    br = item.sceneBoundingRect()
                    visible_rect = br if visible_rect.isEmpty() else visible_rect.united(br)

        if not visible_rect.isEmpty():
            visible_rect.adjust(-35, -35, 35, 35)
            self.circuit_view.fitInView(visible_rect, Qt.KeepAspectRatio)
            # Dériver le pourcentage réel du zoom
            scale_x = self.circuit_view.transform().m11()
            pct = int(round(scale_x * 100))
            self.zoom_slider.blockSignals(True)
            self.zoom_slider.setValue(max(50, min(200, pct)))
            self.zoom_val_label.setText(f"{max(50, min(200, pct))}%")
            self.zoom_slider.blockSignals(False)
        else:
            self.circuit_view.resetTransform()
            self.circuit_view.centerOn(220, 160)
            self.zoom_slider.setValue(100)

    def _sync_scene_to_project(self):
        """Synchronise l'état physique exact de la scène et des cavaliers vers le modèle du projet."""
        self.current_project.connections = list(self.circuit_scene.connections)
        for comp in self.current_project.components:
            item = self.circuit_scene.component_items.get(comp.id)
            if item:
                comp.x = item.scenePos().x()
                comp.y = item.scenePos().y()
                comp.rotation = item.rotation()
                holes = self.circuit_scene.topology.get_component_holes(comp.id)
                comp.pin_insertions = dict(holes)

    def _mark_as_clean(self):
        """Réinitialise tous les indicateurs de modification du projet."""
        self._is_manually_modified = False
        if hasattr(self, "circuit_scene") and self.circuit_scene:
            self.circuit_scene.is_modified = False
        if hasattr(self, "undo_stack") and self.undo_stack:
            self.undo_stack.setClean()
        if hasattr(self, "code_editor") and self.code_editor and hasattr(self.code_editor, "document"):
            self.code_editor.document().setModified(False)

    def is_project_modified(self) -> bool:
        """Indique si le projet actuel a été modifié depuis le dernier enregistrement."""
        if getattr(self, "_is_manually_modified", False):
            return True
        if hasattr(self, "circuit_scene") and getattr(self.circuit_scene, "is_modified", False):
            return True
        if hasattr(self, "undo_stack") and self.undo_stack and not self.undo_stack.isClean():
            return True
        if hasattr(self, "code_editor") and self.code_editor and hasattr(self.code_editor, "document"):
            if self.code_editor.document().isModified():
                return True
        return False

    def _maybe_save_changes(self, action_name: str = "continuer") -> bool:
        """Demande confirmation à l'utilisateur s'il y a des modifications non enregistrées.
        Retourne True si l'opération peut se poursuivre, False si elle doit être annulée."""
        if not self.is_project_modified():
            return True

        # En mode test automatisé (pytest), autoriser la continuation sans bloquer,
        # sauf si le test active explicitement le test du dialogue.
        import os
        if os.environ.get("PYTEST_CURRENT_TEST") and not getattr(self, "_test_save_confirmation", False):
            return True

        from ..app.i18n import tr
        dlg = QMessageBox(self)
        dlg.setWindowTitle(tr("save_changes_title"))
        dlg.setText(tr("save_changes_msg").replace("{action}", action_name))
        dlg.setIcon(QMessageBox.Question)

        btn_save = dlg.addButton(tr("btn_save"), QMessageBox.AcceptRole)
        btn_discard = dlg.addButton(tr("btn_dont_save"), QMessageBox.DestructiveRole)
        btn_cancel = dlg.addButton(tr("btn_cancel"), QMessageBox.RejectRole)
        dlg.setDefaultButton(btn_save)

        dlg.exec()
        clicked = dlg.clickedButton()
        if clicked == btn_save:
            return bool(self._save_project())
        elif clicked == btn_discard:
            return True
        else:
            return False

    def _new_project(self):
        if not self._maybe_save_changes("créer un nouveau montage"):
            return
        self.current_project = ProjectService.create_empty_project()
        self.current_file_path = None
        self._load_project(self.current_project)
        self.setWindowTitle("ESP32 MicroPython Lab")
        self._show_workspace_view()
        self.status_bar.showMessage("Nouveau montage vierge créé (laboratoire prêt pour votre montage)", 4000)

    def _close_project(self):
        """Ferme le projet en cours et revient à la page d'accueil."""
        if not self._maybe_save_changes("fermer le projet"):
            return
        if hasattr(self, "simulation_engine") and (
            (callable(getattr(self.simulation_engine, "is_running", None)) and self.simulation_engine.is_running())
            or getattr(self.simulation_engine, "is_running", False) is True
        ):
            self._stop_simulation()
        self.current_project = ProjectService.create_empty_project()
        self.current_file_path = None
        self._load_project(self.current_project)
        self._show_welcome_view()
        self.setWindowTitle("ESP32 MicroPython Lab — Coder · Simuler · Expérimenter · Apprendre")
        self.status_bar.showMessage("Projet fermé. Page d'accueil active.", 3000)

    def _save_project(self) -> bool:
        if getattr(self, "app_profile", "teacher") == "student":
            if not self.current_project.pedagogy_profile or "student_submission" not in self.current_project.pedagogy_profile:
                if self.current_project.pedagogy_profile and self.current_project.pedagogy_profile.get("is_exam"):
                    return self._submit_student_work()
        self.current_project.set_main_code(self.code_editor.toPlainText())
        self._sync_scene_to_project()
        
        if not self.current_file_path:
            return self._save_project_as()
            
        self.current_project.name = self.current_file_path.stem

        if getattr(self, "app_profile", "teacher") == "teacher":
            ProjectService.save_project(self.current_project, self.current_file_path)
            self.setWindowTitle(f"ESP32 MicroPython Lab — {self.current_file_path.name}")
            self.status_bar.showMessage(f"Projet sauvegardé : {self.current_file_path.name}", 3000)
            self._add_recent_file(str(self.current_file_path))
            self._mark_as_clean()
            return True
        else:
            return bool(self._save_project_as())

    def _save_project_as(self) -> bool:
        if getattr(self, "app_profile", "teacher") == "student":
            if not self.current_project.pedagogy_profile or "student_submission" not in self.current_project.pedagogy_profile:
                if self.current_project.pedagogy_profile and self.current_project.pedagogy_profile.get("is_exam"):
                    return self._submit_student_work()
        self.current_project.set_main_code(self.code_editor.toPlainText())
        self._sync_scene_to_project()
        
        default_name = ""
        if getattr(self, "app_profile", "teacher") == "student" and self.current_project.pedagogy_profile and "student_submission" in self.current_project.pedagogy_profile:
            sub = self.current_project.pedagogy_profile["student_submission"]
            nom = sub.get("nom", "").strip()
            prenom = sub.get("prenom", "").strip()
            classe = sub.get("classe", "").strip()
            act_title = self.current_project.pedagogy_profile.get("activity", {}).get("title", "TP")
            act_title = act_title.replace(" ", "_").replace("/", "-")
            if nom and prenom:
                default_name = f"{nom}_{prenom}_{classe}_{act_title}.lab32"

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Sauvegarder le projet sous...",
            default_name,
            "Projet ESP32 Lab (*.lab32);;Archive ESP32 Lab (*.esp32lab);;Fichier JSON (*.json);;Tous les fichiers (*.*)"
        )
        if path:
            if not any(path.lower().endswith(ext) for ext in [".lab32", ".esp32lab", ".json"]):
                path += ".lab32"
            self.current_file_path = Path(path)
            self.current_project.name = self.current_file_path.stem

        if getattr(self, "app_profile", "teacher") == "teacher":
            ProjectService.save_project(self.current_project, self.current_file_path)
            self.setWindowTitle(f"ESP32 MicroPython Lab — {self.current_file_path.name}")
            self.status_bar.showMessage(f"Projet enregistré sous {self.current_file_path.name}", 3000)
            self._add_recent_file(str(self.current_file_path))
            self._mark_as_clean()
            return True
        return False

    def _open_project(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Ouvrir un projet",
            "",
            "Projet ESP32 Lab (*.lab32 *.esp32lab *.json);;Tous les fichiers (*.*)"
        )
        if path:
            self.open_project_file(path)

    def open_project_file(self, path: str | Path, check_dirty: bool = True) -> bool:
        """Ouvre un fichier projet (.lab32, .esp32lab, .json) directement."""
        if check_dirty and not self._maybe_save_changes("ouvrir un autre projet"):
            return False
        try:
            p = Path(path)
            if not p.exists():
                return False
            project = ProjectService.load_project(p)
            self.current_file_path = p
            project.name = p.stem
            self._load_project(project)
            self.setWindowTitle(f"ESP32 MicroPython Lab — {self.current_file_path.name}")
            self._show_workspace_view()
            self.status_bar.showMessage(f"Projet chargé : {self.current_file_path.name}", 3000)
            self._add_recent_file(str(p))
            return True
        except Exception as ex:
            QMessageBox.critical(self, "Erreur d'ouverture", f"Impossible d'ouvrir le projet :\n{ex}")
            return False

    def _add_recent_file(self, file_path: str):
        settings = QSettings("ESP32Lab", "ESP32MicroPythonLab")
        files = settings.value("recent_files", [])
        if not isinstance(files, list):
            files = []
        if file_path in files:
            files.remove(file_path)
        files.insert(0, file_path)
        if len(files) > 5:
            files = files[:5]
        settings.setValue("recent_files", files)
        self._update_recent_files_menu()

    def _update_recent_files_menu(self):
        menus = []
        if hasattr(self, "recent_menu") and self.recent_menu:
            menus.append(self.recent_menu)
        if hasattr(self, "menu_recent") and self.menu_recent:
            menus.append(self.menu_recent)

        for m in menus:
            m.clear()

        settings = QSettings("ESP32Lab", "ESP32MicroPythonLab")
        files = settings.value("recent_files", [])
        if not isinstance(files, list):
            files = []

        from ..app.i18n import tr
        if not files:
            for m in menus:
                act = QAction(tr("no_recent_files"), self)
                act.setEnabled(False)
                m.addAction(act)
        else:
            for file_path in files:
                for m in menus:
                    act = QAction(file_path, self)
                    act.triggered.connect(lambda checked=False, p=file_path: self._open_recent_file(p))
                    m.addAction(act)

    def _open_recent_file(self, path: str):
        self.open_project_file(path)

    # =========================================================================
    # Barre de menus native complète (Fichier, Édition, Composants, Affichage,
    # Simulation, Cours & TPs, Exemples, Aide)
    # =========================================================================

    def _create_menu_bar(self):
        from ..app.i18n import tr, get_current_language

        # Détacher les actions d'annulation/rétablissement existantes pour éviter les doublons
        if hasattr(self, "act_undo") and self.act_undo in self.actions():
            self.removeAction(self.act_undo)
        if hasattr(self, "act_redo") and self.act_redo in self.actions():
            self.removeAction(self.act_redo)

        menubar = self.menuBar()
        menubar.clear()

        # -----------------------------------------------------------------
        # 1. 📁 FICHIER / FILE
        # -----------------------------------------------------------------
        menu_file = menubar.addMenu(tr("menu_file"))

        act_home = QAction(tr("act_home"), self)
        act_home.setShortcut(QKeySequence("Ctrl+H"))
        act_home.triggered.connect(self._show_welcome_view)
        menu_file.addAction(act_home)

        menu_file.addSeparator()

        act_new = QAction(tr("act_new"), self)
        act_new.setShortcut(QKeySequence.New)
        act_new.triggered.connect(self._new_project)
        menu_file.addAction(act_new)

        act_open = QAction(tr("act_open"), self)
        act_open.setShortcut(QKeySequence.Open)
        act_open.triggered.connect(self._open_project)
        menu_file.addAction(act_open)

        self.menu_recent = menu_file.addMenu(tr("menu_recent"))
        self._update_recent_files_menu()

        act_close = QAction(tr("act_close"), self)
        act_close.setShortcut(QKeySequence("Ctrl+W"))
        act_close.triggered.connect(self._close_project)
        menu_file.addAction(act_close)

        menu_file.addSeparator()

        act_save = QAction(tr("act_save"), self)
        act_save.setShortcut(QKeySequence.Save)
        act_save.triggered.connect(self._save_project)
        menu_file.addAction(act_save)

        act_save_as = QAction(tr("act_save_as"), self)
        act_save_as.setShortcut(QKeySequence.SaveAs)
        act_save_as.triggered.connect(self._save_project_as)
        menu_file.addAction(act_save_as)

        menu_file.addSeparator()

        act_export_report = QAction(tr("act_export_report"), self)
        act_export_report.triggered.connect(self._export_lab_report)
        menu_file.addAction(act_export_report)

        act_export_bom = QAction(tr("act_export_bom"), self)
        act_export_bom.triggered.connect(self._export_bom_csv)
        menu_file.addAction(act_export_bom)

        menu_file.addSeparator()

        act_quit = QAction(tr("act_exit"), self)
        act_quit.setShortcut(QKeySequence.Quit)
        act_quit.triggered.connect(self.close)
        menu_file.addAction(act_quit)

        # -----------------------------------------------------------------
        # 2. ✏️ ÉDITION / EDIT
        # -----------------------------------------------------------------
        menu_edit = menubar.addMenu(tr("menu_edit"))

        self.act_undo = self.undo_stack.createUndoAction(self, tr("act_undo"))
        self.act_undo.setShortcut(QKeySequence.Undo)
        menu_edit.addAction(self.act_undo)
        self.addAction(self.act_undo)

        self.act_redo = self.undo_stack.createRedoAction(self, tr("act_redo"))
        self.act_redo.setShortcut(QKeySequence.Redo)
        menu_edit.addAction(self.act_redo)
        self.addAction(self.act_redo)

        menu_edit.addSeparator()

        act_rotate = QAction(tr("act_rotate"), self)
        act_rotate.setShortcut(QKeySequence(Qt.Key_R))
        act_rotate.triggered.connect(lambda: self.circuit_scene.rotate_selected_component(90.0))
        menu_edit.addAction(act_rotate)

        act_delete = QAction(tr("act_delete"), self)
        act_delete.setShortcut(QKeySequence(Qt.Key_Delete))
        act_delete.triggered.connect(self.circuit_scene.delete_selected_items)
        menu_edit.addAction(act_delete)

        menu_edit.addSeparator()

        act_select_all = QAction(tr("act_select_all"), self)
        act_select_all.setShortcut(QKeySequence.SelectAll)
        act_select_all.triggered.connect(self._select_all_items)
        menu_edit.addAction(act_select_all)

        act_clear_circuit = QAction(tr("act_clear_circuit"), self)
        act_clear_circuit.triggered.connect(self._clear_circuit_workspace)
        menu_edit.addAction(act_clear_circuit)

        # -----------------------------------------------------------------
        # 3. 🧩 COMPOSANTS / COMPONENTS
        # -----------------------------------------------------------------
        menu_comp = menubar.addMenu(tr("menu_comp"))

        # Cartes
        menu_boards = menu_comp.addMenu(tr("menu_boards"))
        act_add_esp32 = QAction(tr("comp_esp32"), self)
        act_add_esp32.triggered.connect(lambda: self._on_add_component("esp32"))
        menu_boards.addAction(act_add_esp32)

        # Basiques
        menu_basics = menu_comp.addMenu(tr("menu_basics"))
        components_basics = [
            ("led", tr("comp_led")),
            ("resistor", tr("comp_resistor")),
            ("button", tr("comp_button")),
            ("switch", tr("comp_switch")),
            ("potentiometer", tr("comp_pot")),
            ("joystick", tr("comp_joystick")),
            ("buzzer", tr("comp_buzzer")),
            ("relay", tr("comp_relay")),
            ("servo", tr("comp_servo")),
        ]
        for ctype, clabel in components_basics:
            act = QAction(clabel, self)
            act.triggered.connect(lambda checked=False, ct=ctype: self._on_add_component(ct))
            menu_basics.addAction(act)

        # Afficheurs
        menu_displays = menu_comp.addMenu(tr("menu_displays"))
        components_displays = [
            ("rgb_led", tr("comp_rgb_led")),
            ("neopixel", tr("comp_neopixel")),
            ("oled", tr("comp_oled")),
            ("lcd", tr("comp_lcd")),
        ]
        for ctype, clabel in components_displays:
            act = QAction(clabel, self)
            act.triggered.connect(lambda checked=False, ct=ctype: self._on_add_component(ct))
            menu_displays.addAction(act)

        # Capteurs
        menu_sensors = menu_comp.addMenu(tr("menu_sensors"))
        components_sensors = [
            ("ldr", tr("comp_ldr")),
            ("pir", tr("comp_pir")),
            ("dht22", tr("comp_dht22")),
            ("hcsr04", tr("comp_hcsr04")),
        ]
        for ctype, clabel in components_sensors:
            act = QAction(clabel, self)
            act.triggered.connect(lambda checked=False, ct=ctype: self._on_add_component(ct))
            menu_sensors.addAction(act)

        # -----------------------------------------------------------------
        # 4. 👁️ AFFICHAGE / VIEW
        # -----------------------------------------------------------------
        menu_view = menubar.addMenu(tr("menu_view"))

        act_v_sim = QAction(tr("act_v_sim"), self)
        act_v_sim.triggered.connect(self._show_simulation_view)
        menu_view.addAction(act_v_sim)

        act_v_schema = QAction(tr("act_v_schema"), self)
        act_v_schema.triggered.connect(self._show_schematic_view)
        menu_view.addAction(act_v_schema)

        menu_view.addSeparator()

        act_zoom_in = QAction(tr("act_zoom_in"), self)
        act_zoom_in.setShortcut(QKeySequence.ZoomIn)
        act_zoom_in.triggered.connect(lambda: self.zoom_slider.setValue(min(200, self.zoom_slider.value() + 10)))
        menu_view.addAction(act_zoom_in)

        act_zoom_out = QAction(tr("act_zoom_out"), self)
        act_zoom_out.setShortcut(QKeySequence.ZoomOut)
        act_zoom_out.triggered.connect(lambda: self.zoom_slider.setValue(max(50, self.zoom_slider.value() - 10)))
        menu_view.addAction(act_zoom_out)

        act_zoom_fit = QAction(tr("act_zoom_fit"), self)
        act_zoom_fit.setShortcut(QKeySequence("Ctrl+0"))
        act_zoom_fit.triggered.connect(self._fit_view)
        menu_view.addAction(act_zoom_fit)

        menu_view.addSeparator()

        act_toggle_montage = QAction(tr("act_toggle_montage"), self)
        act_toggle_montage.setCheckable(True)
        act_toggle_montage.setChecked(True)
        act_toggle_montage.toggled.connect(self.simulator_widget.setVisible)
        menu_view.addAction(act_toggle_montage)

        act_toggle_palette = QAction(tr("act_toggle_palette"), self)
        act_toggle_palette.setCheckable(True)
        act_toggle_palette.setChecked(True)
        act_toggle_palette.toggled.connect(self.palette_widget.setVisible)
        menu_view.addAction(act_toggle_palette)

        act_toggle_props = QAction(tr("act_toggle_props"), self)
        act_toggle_props.setCheckable(True)
        act_toggle_props.setChecked(True)
        act_toggle_props.toggled.connect(self.properties_panel.setVisible)
        menu_view.addAction(act_toggle_props)

        act_toggle_console = QAction(tr("act_toggle_console"), self)
        act_toggle_console.triggered.connect(lambda: self.bottom_tabs.setCurrentIndex(0))
        menu_view.addAction(act_toggle_console)

        act_toggle_scope = QAction(tr("act_toggle_scope"), self)
        act_toggle_scope.triggered.connect(lambda: self.bottom_tabs.setCurrentIndex(1))
        menu_view.addAction(act_toggle_scope)

        act_toggle_erc = QAction(tr("act_toggle_erc"), self)
        act_toggle_erc.triggered.connect(lambda: self.bottom_tabs.setCurrentIndex(2))
        menu_view.addAction(act_toggle_erc)

        menu_view.addSeparator()

        # Sous-menu Thème (Sombre / Clair)
        menu_theme = menu_view.addMenu(tr("menu_theme"))
        theme_group = QActionGroup(self)
        theme_group.setExclusive(True)

        settings = QSettings("ESP32Lab", "ESP32MicroPythonLab")
        curr_theme = settings.value("theme", "dark")

        self.act_theme_dark = QAction(tr("theme_dark"), self)
        self.act_theme_dark.setCheckable(True)
        self.act_theme_dark.setChecked(curr_theme == "dark")
        self.act_theme_dark.triggered.connect(lambda: self._set_app_theme("dark"))
        theme_group.addAction(self.act_theme_dark)
        menu_theme.addAction(self.act_theme_dark)

        self.act_theme_light = QAction(tr("theme_light"), self)
        self.act_theme_light.setCheckable(True)
        self.act_theme_light.setChecked(curr_theme == "light")
        self.act_theme_light.triggered.connect(lambda: self._set_app_theme("light"))
        theme_group.addAction(self.act_theme_light)
        menu_theme.addAction(self.act_theme_light)

        # Sous-menu Langue (Français / English)
        menu_lang = menu_view.addMenu(tr("menu_language"))
        lang_group = QActionGroup(self)
        lang_group.setExclusive(True)

        curr_lang = get_current_language()

        self.act_lang_fr = QAction("🇫🇷 Français", self)
        self.act_lang_fr.setCheckable(True)
        self.act_lang_fr.setChecked(curr_lang == "fr")
        self.act_lang_fr.triggered.connect(lambda: self._set_app_language("fr"))
        lang_group.addAction(self.act_lang_fr)
        menu_lang.addAction(self.act_lang_fr)

        self.act_lang_en = QAction("🇬🇧 English", self)
        self.act_lang_en.setCheckable(True)
        self.act_lang_en.setChecked(curr_lang == "en")
        self.act_lang_en.triggered.connect(lambda: self._set_app_language("en"))
        lang_group.addAction(self.act_lang_en)
        menu_lang.addAction(self.act_lang_en)

        # -----------------------------------------------------------------
        # 5. ⚡ SIMULATION
        # -----------------------------------------------------------------
        menu_sim = menubar.addMenu(tr("menu_sim"))

        act_run = QAction(tr("act_run"), self)
        act_run.setShortcut(Qt.Key_F5)
        act_run.triggered.connect(self._start_simulation)
        menu_sim.addAction(act_run)

        act_stop = QAction(tr("act_stop"), self)
        act_stop.setShortcut(Qt.Key_F6)
        act_stop.triggered.connect(self._stop_simulation)
        menu_sim.addAction(act_stop)

        act_reset = QAction(tr("act_reset"), self)
        act_reset.triggered.connect(self._reset_simulation)
        menu_sim.addAction(act_reset)

        act_check_erc = QAction(tr("act_erc"), self)
        act_check_erc.setShortcut(QKeySequence("F7"))
        act_check_erc.triggered.connect(self._run_erc_check)
        menu_sim.addAction(act_check_erc)

        menu_sim.addSeparator()

        act_upload = QAction(tr("act_upload"), self)
        act_upload.triggered.connect(self._upload_to_esp32)
        menu_sim.addAction(act_upload)

        menu_sim.addSeparator()

        # Passerelle Réseau Réelle (Host Network Bridge)
        from ..simulator.modules.network import is_host_bridge_enabled, set_host_bridge_enabled
        saved_bridge = settings.value("use_host_network_bridge", True, type=bool)
        set_host_bridge_enabled(saved_bridge)

        self.act_network_bridge = QAction(tr("act_bridge"), self)
        self.act_network_bridge.setCheckable(True)
        self.act_network_bridge.setChecked(saved_bridge)
        self.act_network_bridge.setToolTip(tr("act_bridge_tip"))
        self.act_network_bridge.toggled.connect(self._on_toggle_network_bridge)
        menu_sim.addAction(self.act_network_bridge)

        
        # -----------------------------------------------------------------
        # 5.5. ACTIVITÉ (Enseignant uniquement)
        # -----------------------------------------------------------------
        if getattr(self, "app_profile", "student") == "teacher":
            menu_activity = menubar.addMenu("Activité")
            
            act_new = QAction("Créer une nouvelle activité...", self)
            act_new.triggered.connect(self._start_activity_wizard)
            menu_activity.addAction(act_new)
            
            self.act_edit_activity = QAction("Modifier l'Énoncé...", self)
            self.act_edit_activity.triggered.connect(self._edit_activity_metadata)
            menu_activity.addAction(self.act_edit_activity)
            
            if hasattr(self, "act_lock_tool"):
                menu_activity.addAction(self.act_lock_tool)
                
            self.act_finish_activity = QAction("Exporter l'Activité...", self)
            self.act_finish_activity.triggered.connect(self._export_activity)
            menu_activity.addAction(self.act_finish_activity)
                
            
        # -----------------------------------------------------------------
        # 6. 🎓 COURS & TPS
        # -----------------------------------------------------------------
        self.menu_courses = menubar.addMenu(tr("menu_courses"))
        menu_courses = self.menu_courses

        act_all_courses = QAction(tr("act_all_courses"), self)
        act_all_courses.triggered.connect(self._show_courses_dialog)
        menu_courses.addAction(act_all_courses)

        act_evaluate = QAction(tr("act_evaluate"), self)
        act_evaluate.triggered.connect(self._evaluate_current_circuit)
        menu_courses.addAction(act_evaluate)

        menu_courses.addSeparator()

        courses_meta = [
            ("lesson_1_gpio_led", "TP 1 : LED Clignotante (Blink)"),
            ("lesson_2_button", "TP 2 : Bouton-Poussoir & Entrée Numérique"),
            ("lesson_3_pot_adc", "TP 3 : Variateur PWM & Potentiomètre"),
            ("lesson_4_servo_pwm", "TP 4 : Pilotage Servomoteur SG90"),
            ("lesson_5_dht22", "TP 5 : Capteur DHT22 & Station Météo"),
            ("lesson_6_oled", "TP 6 : Écran OLED SSD1306 (I2C)"),
            ("lesson_7_hcsr04", "TP 7 : Mesure de distance Ultrasons"),
            ("lesson_8_lcd", "TP 8 : Écran LCD 1602 (I2C)"),
            ("lesson_9_relay", "TP 9 : Domotique & Relais 5V"),
            ("lesson_10_neopixel", "TP 10 : Ruban Adressable NeoPixel"),
        ]
        for cid, clabel in courses_meta:
            act = QAction(clabel, self)
            act.triggered.connect(lambda checked=False, course_id=cid: self._load_course_by_id(course_id))
            menu_courses.addAction(act)

        # -----------------------------------------------------------------
        # 7. 💡 EXEMPLES
        # -----------------------------------------------------------------
        self.menu_examples = menubar.addMenu(tr("menu_examples"))
        menu_examples = self.menu_examples

        act_all_examples = QAction(tr("act_all_examples"), self)
        act_all_examples.triggered.connect(self._show_examples_dialog)
        menu_examples.addAction(act_all_examples)

        menu_examples.addSeparator()

        examples_quick = [
            ("blink", "💡 1. Clignotement LED (Blink)"),
            ("button_led", "🔘 2. Bouton Poussoir & LED"),
            ("pot_pwm", "🎛️ 3. Variateur Potentiomètre & PWM"),
            ("servo_sweep", "🦾 4. Balayage Servomoteur SG90"),
            ("dht22_weather", "🌡️ 5. Station Météo DHT22"),
            ("oled_display", "📱 6. Écran OLED SSD1306 (I2C)"),
            ("neopixel_ring", "🌈 7. Anneau LED NeoPixel WS2812"),
            ("relay_control", "🔀 8. Commande Relais 5V"),
            ("weather_cloud_iot", "🌦️ 9. Station Météo Cloud IoT"),
            ("wifi_scanner", "🔍 10. Scanner Wi-Fi Réel"),
            ("mqtt_client", "🌐 11. Client MQTT IoT (HiveMQ)"),
            ("ble_scanner", "📱 12. Scanner Bluetooth BLE"),
        ]
        for ekey, elabel in examples_quick:
            act = QAction(elabel, self)
            act.triggered.connect(lambda checked=False, k=ekey: self._load_example_by_key(k))
            menu_examples.addAction(act)

        # -----------------------------------------------------------------
        # 8. ❓ AIDE
        # -----------------------------------------------------------------
        menu_help = menubar.addMenu(tr("menu_help"))

        act_guide = QAction(tr("act_guide"), self)
        act_guide.triggered.connect(self._show_breadboard_help_dialog)
        menu_help.addAction(act_guide)

        act_shortcuts = QAction(tr("act_shortcuts"), self)
        act_shortcuts.triggered.connect(self._show_shortcuts_dialog)
        menu_help.addAction(act_shortcuts)

        menu_help.addSeparator()

        self.act_updates = QAction(tr("act_updates"), self)
        self.act_updates.triggered.connect(lambda: self._check_for_updates(manual=True))
        menu_help.addAction(self.act_updates)

        act_about = QAction(tr("act_about"), self)
        act_about.triggered.connect(self._show_about_dialog)
        menu_help.addAction(act_about)

    def _select_all_items(self):
        if self.code_editor.hasFocus():
            self.code_editor.selectAll()
        else:
            bb = self.circuit_scene.breadboard_item
            for item in self.circuit_scene.items():
                if hasattr(item, "component_id") and item != bb and getattr(item, "component_id", "") not in ("breadboard_1", "esp32", "board"):
                    item.setSelected(True)
                elif isinstance(item, WireGraphicsItem):
                    item.setSelected(True)

    def _clear_circuit_workspace(self):
        reply = QMessageBox.question(
            self,
            "Effacer le montage",
            "Voulez-vous supprimer tous les composants et fils du circuit actuel ?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )
        if reply == QMessageBox.Yes:
            bb = self.circuit_scene.breadboard_item
            to_delete = [
                it for it in self.circuit_scene.items()
                if (hasattr(it, "component_id") and it != bb and getattr(it, "component_id", "") not in ("breadboard_1", "esp32", "board"))
                or isinstance(it, WireGraphicsItem)
            ]
            if to_delete:
                if self.undo_stack:
                    from .canvas.undo_commands import DeleteItemsCommand
                    self.undo_stack.push(DeleteItemsCommand(self.circuit_scene, to_delete))
                else:
                    self.circuit_scene._execute_delete_items(to_delete)

    def _load_course_by_id(self, course_id: str):
        from ..core.courses import get_course_curriculum
        lessons = get_course_curriculum()
        cid = course_id.lower().strip()
        target = next((l for l in lessons if l.id.lower() == cid), None)
        if not target:
            # Fallback pour raccourcis comme 'tp1' -> 'lesson_1_...'
            num = "".join(filter(str.isdigit, cid))
            for l in lessons:
                if cid in l.id.lower() or (num and f"lesson_{num}_" in l.id.lower()):
                    target = l
                    break
        if target and target.starter_project:
            self._current_lesson = target
            self._load_project(target.starter_project)
            self._show_workspace_view()
            self.status_bar.showMessage(f"Exercice chargé : {target.title}", 4000)

    def _evaluate_current_circuit(self):
        from ..core.evaluator import ExerciseEvaluator
        from ..core.courses import get_course_curriculum
        self.current_project.set_main_code(self.code_editor.toPlainText())
        self._sync_scene_to_project()

        lesson = getattr(self, "_current_lesson", None)
        if not lesson:
            lessons = get_course_curriculum()
            lesson = lessons[0] if lessons else None

        if not lesson:
            QMessageBox.information(self, "Auto-évaluation", "Aucune leçon disponible pour l'évaluation.")
            return

        evaluator = ExerciseEvaluator()
        res = evaluator.evaluate(lesson, self.current_project)

        status_str = "Validé avec succès ✅" if res.passed else "À corriger ❌"
        msg = f"Auto-évaluation : {res.lesson_title}\n"
        msg += f"Score : {res.score} / {res.max_score} ({res.percentage}%)\n"
        msg += f"Statut : {status_str}\n\n"
        msg += "Détails des critères :\n"
        for c in res.criteria:
            icon = "✅" if c.passed else "❌"
            feedback_txt = f" : {c.feedback}" if c.feedback else ""
            msg += f"{icon} {c.title} ({c.points}/{c.max_points} pts){feedback_txt}\n"
        if res.general_feedback:
            msg += f"\nConseil : {res.general_feedback}"

        QMessageBox.information(self, "Auto-évaluation pédagogique", msg)

    def _load_example_by_key(self, key: str):
        from ..core.examples import get_all_examples
        examples = get_all_examples()
        if key in examples:
            self._load_project(examples[key])
            self._show_workspace_view()
            self.status_bar.showMessage(f"Exemple chargé : {examples[key].name}", 4000)

    def _export_lab_report(self):
        from PySide6.QtWidgets import QFileDialog, QMessageBox, QDialog, QVBoxLayout, QLabel, QComboBox, QDialogButtonBox
        from PySide6.QtGui import QTextDocument, QPageSize, QPageLayout, QPainter, QImage
        from PySide6.QtPrintSupport import QPrinter
        import datetime
        
        dialog = QDialog(self)
        dialog.setWindowTitle("Format d'exportation")
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel("Choisissez le format :"))
        combo = QComboBox()
        combo.addItems(["Document Word (*.docx)", "Document PDF (*.pdf)"])
        layout.addWidget(combo)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dialog.accept)
        btns.rejected.connect(dialog.reject)
        layout.addWidget(btns)
        
        if dialog.exec() != QDialog.Accepted:
            return
            
        is_pdf = (combo.currentIndex() == 1)
        ext = ".pdf" if is_pdf else ".docx"
        filter_str = "Fichiers PDF (*.pdf)" if is_pdf else "Documents Word (*.docx)"
        
        filepath, _ = QFileDialog.getSaveFileName(self, "Exporter le compte rendu", f"Projet_{self.current_project.name}{ext}", filter_str)
        if not filepath:
            return
            
        try:
            # Capture circuit image
            import base64
            from PySide6.QtCore import QByteArray, QBuffer, QRectF
            
            rect = self.circuit_scene.itemsBoundingRect()
            if rect.isEmpty():
                rect = QRectF(0, 0, 800, 600)
            
            img = QImage(rect.size().toSize(), QImage.Format_ARGB32_Premultiplied)
            img.fill(0xffffffff) # white background
            painter = QPainter(img)
            self.circuit_scene.render(painter, QRectF(img.rect()), rect)
            painter.end()
            
            img = img.scaledToWidth(800)
            
            ba = QByteArray()
            buffer = QBuffer(ba)
            buffer.open(QBuffer.WriteOnly)
            img.save(buffer, "PNG")
            img_b64 = ba.toBase64().data().decode()
            
            now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
            code_text = self.code_editor.toPlainText()
            
            if is_pdf:
                html = f"""
                <html>
                <head><style>
                    body {{ font-family: sans-serif; font-size: 11pt; }}
                    h1 {{ color: #2c3e50; border-bottom: 1px solid #ccc; }}
                    h2 {{ color: #34495e; margin-top: 20pt; }}
                    pre {{ background-color: #f8f9fa; padding: 10pt; border: 1px solid #ddd; }}
                    .footer {{ font-size: 9pt; color: #7f8c8d; text-align: right; margin-top: 30pt; }}
                </style></head>
                <body>
                <h1>ESP32 MicroPython Lab - Compte Rendu</h1>
                <p><b>Projet :</b> {self.current_project.name}<br/><b>Date :</b> {now_str}</p>
                <h2>Montage / Circuit</h2>
                <img src="data:image/png;base64,{img_b64}" width="100%" />
                <h2>Code MicroPython</h2>
                <pre>{code_text.replace('<','&lt;').replace('>','&gt;')}</pre>
                <div class="footer">Logiciel ESP32 MicroPython Lab - Généré automatiquement</div>
                </body></html>
                """
                doc = QTextDocument()
                doc.setHtml(html)
                printer = QPrinter(QPrinter.PrinterResolution)
                printer.setOutputFormat(QPrinter.PdfFormat)
                printer.setOutputFileName(filepath)
                
                layout = QPageLayout()
                layout.setPageSize(QPageSize(QPageSize.A4))
                printer.setPageLayout(layout)
                
                doc.print_(printer)
            else:
                try:
                    import docx
                    from docx.shared import Pt, Inches
                    import io
                except ImportError:
                    QMessageBox.warning(self, "Erreur", "Le module 'python-docx' est requis pour exporter au format Word.\nInstallez-le avec 'pip install python-docx'.")
                    return
                
                doc = docx.Document()
                
                # Header
                section = doc.sections[0]
                header = section.header
                hp = header.paragraphs[0]
                hp.text = f"ESP32 MicroPython Lab - {self.current_project.name}"
                hp.style.font.size = Pt(9)
                hp.style.font.color.rgb = docx.shared.RGBColor(127,140,141)
                
                # Footer
                footer = section.footer
                fp = footer.paragraphs[0]
                fp.text = f"Généré le {now_str}"
                fp.style.font.size = Pt(9)
                fp.style.font.color.rgb = docx.shared.RGBColor(127,140,141)
                
                doc.add_heading("ESP32 MicroPython Lab - Compte Rendu", 0)
                doc.add_paragraph(f"Projet : {self.current_project.name}\nDate : {now_str}")
                
                doc.add_heading("Montage / Circuit", level=1)
                img_io = io.BytesIO(ba.data())
                doc.add_picture(img_io, width=Inches(6.0))
                
                doc.add_heading("Code MicroPython", level=1)
                p_code = doc.add_paragraph(code_text)
                p_code.style.font.name = 'Courier New'
                p_code.style.font.size = Pt(10)
                
                doc.save(filepath)
                
            QMessageBox.information(self, "Succès", f"Compte rendu exporté avec succès :\n{filepath}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Erreur lors de l'exportation : {e}")
    def _export_bom_csv(self):
        from ..core.bom_exporter import BOMExporter
        self._sync_scene_to_project()
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Exporter la nomenclature des composants (BOM)",
            f"Nomenclature_BOM_{self.current_project.name}.csv",
            "Fichier CSV (*.csv);;Tous les fichiers (*.*)"
        )
        if path:
            exporter = BOMExporter()
            exporter.export_csv(self.current_project, Path(path))
            self.status_bar.showMessage(f"Nomenclature BOM exportée avec succès : {Path(path).name}", 4000)
            QMessageBox.information(
                self,
                "Exportation réussie",
                f"La liste de matériel (BOM) a été enregistrée :\n{path}\n\n"
                "Le fichier est compatible avec Excel, LibreOffice et les outils de CAO."
            )

    def _show_breadboard_help_dialog(self):
        dlg = QMessageBox(self)
        dlg.setWindowTitle("Guide de câblage sur Platine d'Expérimentation (Breadboard MB-102)")
        dlg.setIcon(QMessageBox.Information)
        dlg.setTextFormat(Qt.RichText)
        dlg.setText("""
        <h3>🔌 Règles fondamentales de la platine MB-102 :</h3>
        <ul>
            <li><b>Lignes horizontales (1 à 30) :</b> Les 5 trous d'une même ligne à gauche (<b>a-b-c-d-e</b>) sont reliés électriquement entre eux. Il en va de même à droite (<b>f-g-h-i-j</b>).</li>
            <li><b>Rainure centrale DIP :</b> Isole complètement la colonne <i>e</i> de la colonne <i>f</i>. C'est ici que l'on insère les circuits intégrés et la carte ESP32 sans court-circuiter leurs broches opposées.</li>
            <li><b>Rails d'alimentation verticaux :</b> Les colonnes <b>+ (rouge)</b> et <b>- (bleu/noir)</b> courent tout le long pour distribuer 3.3V / 5V et la Masse (GND).</li>
            <li><b>Fils de liaison Dupont :</b> Cliquez sur une broche pour tirer un fil, ou <b>glissez un embout de fil</b> existant directement d'un trou à un autre avec aimantation !</li>
            <li><b>Insertion des composants :</b> Glissez un composant sur la platine. Le contour vert indique un alignement valide sur les alvéoles.</li>
        </ul>
        """)
        dlg.exec()

    def _show_shortcuts_dialog(self):
        dlg = QMessageBox(self)
        dlg.setWindowTitle("Raccourcis clavier & Commandes interactives")
        dlg.setIcon(QMessageBox.Information)
        dlg.setTextFormat(Qt.RichText)
        dlg.setText("""
        <h3>⌨ Raccourcis et commandes :</h3>
        <table border="0" cellpadding="4">
            <tr><td><b>Ctrl + Z</b></td><td>Annuler la dernière action (déplacement, ajout, suppression, rotation, câblage)</td></tr>
            <tr><td><b>Ctrl + Y / Ctrl+Shift+Z</b></td><td>Rétablir l'action annulée</td></tr>
            <tr><td><b>R</b></td><td>Faire pivoter le composant sélectionné de 90°</td></tr>
            <tr><td><b>Suppr / Backspace</b></td><td>Supprimer le composant ou fil sélectionné</td></tr>
            <tr><td><b>F5</b></td><td>Démarrer la simulation MicroPython</td></tr>
            <tr><td><b>F6</b></td><td>Arrêter la simulation</td></tr>
            <tr><td><b>Ctrl + S</b></td><td>Enregistrer le projet</td></tr>
            <tr><td><b>Ctrl + O</b></td><td>Ouvrir un projet</td></tr>
            <tr><td><b>Ctrl + N</b></td><td>Nouveau montage vierge</td></tr>
            <tr><td><b>Glisser un embout Dupont</b></td><td>Déplacer librement le fil d'un trou à un autre avec aimantation</td></tr>
            <tr><td><b>Molette souris</b></td><td>Zoom avant / arrière</td></tr>
            <tr><td><b>Clic droit glissé / Molette</b></td><td>Panoramique de la vue</td></tr>
        </table>
        """)
        dlg.exec()

    def _show_about_dialog(self):
        from .dialogs.about_dialog import AboutDialog
        dlg = AboutDialog(self)
        dlg.exec()

    def _check_for_updates(self, manual: bool = True):
        from ..core.updater import UpdateCheckerThread
        self.status_bar.showMessage(tr("update_checking"), 3000)
        self._update_thread = UpdateCheckerThread(self)
        self._update_thread.check_finished.connect(lambda res: self._on_update_finished(res, manual))
        self._update_thread.check_failed.connect(lambda err: self._on_update_failed(err, manual))
        self._update_thread.start()

    def _on_update_finished(self, res: dict, manual: bool):
        from PySide6.QtGui import QDesktopServices
        from PySide6.QtCore import QUrl
        if res.get("has_update"):
            latest = res.get("latest_version", "")
            current = res.get("current_version", "")
            notes = res.get("release_notes", "")
            download_url = res.get("download_url", res.get("html_url", ""))

            msg = (
                f"<h3>✨ {tr('update_available_title')}</h3>"
                f"<p>Une nouvelle version <b>{latest}</b> est disponible !<br>"
                f"(Version actuelle installée : <i>{current}</i>)</p>"
            )
            if notes:
                msg += f"<p><b>Notes de version :</b><br><pre style='background:#1e293b; color:#e2e8f0; padding:8px; border-radius:4px;'>{notes[:400]}</pre></p>"

            dlg = QMessageBox(self)
            dlg.setWindowTitle("Mise à jour disponible — ESP32 MicroPython Lab")
            dlg.setIcon(QMessageBox.Information)
            dlg.setTextFormat(Qt.RichText)
            dlg.setText(msg)
            btn_dl = dlg.addButton(tr("update_btn_download"), QMessageBox.AcceptRole)
            btn_cancel = dlg.addButton(tr("update_btn_later"), QMessageBox.RejectRole)
            dlg.exec()

            if dlg.clickedButton() == btn_dl:
                QDesktopServices.openUrl(QUrl(download_url))
        elif manual:
            QMessageBox.information(
                self,
                "Mises à jour — ESP32 MicroPython Lab",
                f"{tr('update_up_to_date')}\n\nVersion actuelle : v{res.get('current_version', '')} ✅"
            )

    def _on_update_failed(self, err_msg: str, manual: bool):
        if manual:
            QMessageBox.warning(
                self,
                "Vérification des mises à jour",
                f"Impossible de vérifier les mises à jour pour le moment :\n{err_msg}"
            )

    def closeEvent(self, event):
        if not self._maybe_save_changes("quitter l'application"):
            event.ignore()
            return
        self._stop_simulation()
        event.accept()

    def _create_new_activity(self):
        if not self._maybe_save_changes("Créer une activité"): return
        from ..core.project_service import ProjectService
        self.current_project = ProjectService.create_empty_project()
        self.current_project.pedagogy_profile = {"is_exam": True}
        self.current_file_path = None
        self._load_project(self.current_project)
        if hasattr(self, "_edit_activity_metadata"): self._edit_activity_metadata()
        
    def _open_student_copy(self):
        if not self._maybe_save_changes("Ouvrir une copie"): return
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        from pathlib import Path
        from ..core.project_service import ProjectService
        path, _ = QFileDialog.getOpenFileName(self, "Ouvrir la copie", "", "Copie (*.lab32)")
        if path:
            try:
                p = ProjectService.load_project(path)
                self.current_project = p
                self.current_file_path = Path(path)
                self._load_project(p)
                self.status_bar.showMessage(f"Copie ouverte : {self.current_file_path.name}", 4000)
            except Exception as e:
                QMessageBox.critical(self, "Erreur", f"Impossible d'ouvrir : {e}")

    def _show_grading_center(self):
        if hasattr(self, "teacher_dashboard"):
            self.central_stack.setCurrentWidget(self.teacher_dashboard)

    def _submit_student_work(self) -> bool:
        from PySide6.QtWidgets import QDialog, QFormLayout, QLineEdit, QLabel, QDialogButtonBox, QMessageBox, QFileDialog
        dialog = QDialog(self)
        dialog.setWindowTitle("Informations Élève")
        layout = QFormLayout(dialog)
        lbl_info = QLabel("Veuillez saisir vos informations pour l'enregistrement de votre TP.")
        lbl_info.setStyleSheet("font-weight: bold; color: #38bdf8; margin-bottom: 10px;")
        layout.addRow(lbl_info)
        
        txt_nom = QLineEdit()
        txt_prenom = QLineEdit()
        txt_classe = QLineEdit()
        
        layout.addRow("Nom :", txt_nom)
        layout.addRow("Prénom :", txt_prenom)
        layout.addRow("Classe :", txt_classe)
        
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dialog.accept)
        btns.rejected.connect(dialog.reject)
        layout.addRow(btns)
        
        if dialog.exec() == QDialog.Accepted:
            nom = txt_nom.text().strip().replace(" ", "_")
            prenom = txt_prenom.text().strip().replace(" ", "_")
            classe = txt_classe.text().strip().replace(" ", "_")
            
            if not nom or not prenom or not classe:
                QMessageBox.warning(self, "Erreur", "Tous les champs sont obligatoires.")
                return False
                
            if not self.current_project.pedagogy_profile:
                self.current_project.pedagogy_profile = {}
                
            self.current_project.pedagogy_profile["student_submission"] = {
                "nom": nom,
                "prenom": prenom,
                "classe": classe,
                "timestamp": __import__('time').time()
            }
            
            act_title = "TP"
            if "activity" in self.current_project.pedagogy_profile:
                act_title = self.current_project.pedagogy_profile["activity"].get("title", "TP")
                act_title = act_title.replace(" ", "_").replace("/", "-")
                
            default_name = f"{act_title}_{nom}_{prenom}_{classe}.lab32"
            
            path, _ = QFileDialog.getSaveFileName(self, "Enregistrer mon travail", default_name, "ESP32 Lab Project (*.lab32)")
            if path:
                self.current_project.set_main_code(self.code_editor.toPlainText())
                self._sync_scene_to_project()
                
                from pathlib import Path
                self.current_file_path = Path(path)
                self.current_project.name = self.current_file_path.stem
        # Generate a new activity ID so it doesn't conflict with the original file's session
        if getattr(self, "app_profile", "teacher") == "teacher":
            if self.current_project.pedagogy_profile and "activity" in self.current_project.pedagogy_profile:
                import uuid
                self.current_project.pedagogy_profile["activity"]["id"] = f"act_{str(uuid.uuid4())[:8]}"

                from ..core.project_service import ProjectService
                ProjectService.save_project(self.current_project, self.current_file_path)
                self.setWindowTitle(f"ESP32 MicroPython Lab - {self.current_file_path.name}")
                self.status_bar.showMessage(f"Projet enregistré : {self.current_file_path.name}", 3000)
                self._add_recent_file(str(self.current_file_path))
                self._mark_as_clean()
                return True
        return False

    def _start_activity_wizard(self):
        from .dialogs.activity_wizard import ActivityWizardDialog
        dlg = ActivityWizardDialog(self)
        if dlg.exec():
            data = dlg.get_activity_data()
            from ..core.project_service import ProjectService
            self.current_project = ProjectService.create_empty_project()
            self.current_project.name = data["title"]
            self.current_project.pedagogy_profile = {
                "activity": {
                    "title": data["title"],
                    "type": data["type"],
                    "duration_minutes": data["duration"],
                    "description": data["description"],
                    "success_criteria": [],
                    "hints": []
                },
                "is_exam": data.get("type", "TP") == "Exam"
            }
            self._show_workspace_view()
            self._load_project(self.current_project)
            self._is_manually_modified = True

    def _on_session_locked(self, event_data=None):
        if getattr(self, "app_profile", "student") == "teacher":
            return
        is_locked = event_data.get("locked", False) if isinstance(event_data, dict) else False
        if hasattr(self, "code_editor"):
            self.code_editor.setReadOnly(is_locked)
        if hasattr(self, "palette_widget"):
            self.palette_widget.setEnabled(not is_locked)

    def _edit_activity_metadata(self):
        from .dialogs.activity_wizard import ActivityWizardDialog
        dlg = ActivityWizardDialog(self)
        if self.current_project and self.current_project.pedagogy_profile:
            act_data = self.current_project.pedagogy_profile.get("activity", {})
            dlg.load_data({
                "title": self.current_project.name,
                "type": act_data.get("type", "TP"),
                "duration": act_data.get("duration_minutes", 60),
                "description": act_data.get("description", "")
            })
        if dlg.exec():
            data = dlg.get_activity_data()
            self.current_project.name = data["title"]
            
            # Preserve other activity stuff
            if not self.current_project.pedagogy_profile:
                self.current_project.pedagogy_profile = {"activity": {}}
                
            self.current_project.pedagogy_profile["is_exam"] = (data.get("type", "TP") == "Exam")
            act = self.current_project.pedagogy_profile.get("activity", {})
            act["title"] = data["title"]
            act["type"] = data["type"]
            act["duration_minutes"] = data["duration"]
            act["description"] = data["description"]
            self.current_project.pedagogy_profile["activity"] = act
            
            if hasattr(self, "activity_panel"):
                self.activity_panel.load_project(self.current_project)
            self._is_manually_modified = True

    def _export_activity(self):
        self.current_project.set_main_code(self.code_editor.toPlainText())
        self._sync_scene_to_project()
        from PySide6.QtWidgets import QFileDialog
        path, _ = QFileDialog.getSaveFileName(
            self, "Exporter l'activité pédagogique",
            f"{self.current_project.name}.lab32",
            "ESP32 Lab Project (*.lab32)"
        )
        if path:
            from esp32_lab.core.project_service import ProjectService
            ProjectService.save_project(self.current_project, path)
            self.status_bar.showMessage(f"Activité exportée avec succès : {path}", 5000)
